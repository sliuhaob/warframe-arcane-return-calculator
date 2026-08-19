#!/usr/bin/env python3
"""Build the Arcane return workbook from the JSON produced by the fetcher."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


SHEET_DETAIL = "满级赋能近48小时平均价"
SHEET_PARAMETERS = "计算参数"
DEFAULT_OUTPUT = Path("outputs") / "daily_arcane_return" / "赋能收益表_最新.xlsx"
ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")

TEAL_DARK = "125C66"
TEAL = "287B84"
TEAL_LIGHT = "D8ECEE"
BODY_TEXT = "22343A"
GRID = "DDE7E9"
WHITE = "FFFFFF"
GREEN_LIGHT = "E2F0D9"
GREEN_TEXT = "2F641C"
YELLOW_LIGHT = "FFF2CC"
YELLOW_TEXT = "9C6500"
RED_LIGHT = "FCE4D6"
RED_TEXT = "C00000"
GREY_LIGHT = "F1F3F4"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="从赋能收益 JSON 生成 Excel 工作簿")
    parser.add_argument("json_path", type=Path, help="抓价程序生成的 JSON 文件")
    parser.add_argument("output_path", nargs="?", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def fallback_valuation_method(row: dict[str, Any], methodology: dict[str, Any]) -> str:
    price = row.get("average_price_48h")
    volume = int(row.get("latest_daily_volume") or 0)
    min_volume = int(methodology.get("minDailyVolume") or 10)
    second_min = int(methodology.get("secondaryFilterMinDailyVolume") or 10)
    second_max = int(methodology.get("secondaryFilterMaxDailyVolume") or 20)
    second_price = float(methodology.get("secondaryFilterMinAveragePrice") or 80)
    recycle = int(row.get("dissolution_vosfor") or 0) > 0
    rejected = (
        price is None
        or volume < min_volume
        or (second_min <= volume <= second_max and float(price) < second_price)
    )
    if rejected:
        return "分解再投" if recycle else "无回收价值"
    return "市场出售"


def _thin_border(color: str = GRID) -> Border:
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def _style_header(row) -> None:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=TEAL_DARK)
        cell.font = Font(name="Microsoft YaHei UI", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border("D6E5E7")


def _write_detail_sheet(
    workbook: Workbook,
    data: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[tuple[str, int, int]]]:
    sheet = workbook.create_sheet(SHEET_DETAIL)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C2"

    summaries = sorted(data["packSummary"], key=lambda row: int(row["rank"]))
    prices = list(data["prices"])
    pack_items = list(data["packItems"])
    methodology = data["methodology"]
    item_by_slug = {row["slug"]: row for row in pack_items}
    rank_by_pack = {row["pack_name_en"]: int(row["rank"]) for row in summaries}
    tier_order = {"普通": 0, "低稀有度": 0, "罕见": 1, "稀有": 2, "传奇": 3}

    def sort_key(row: dict[str, Any]) -> tuple[Any, ...]:
        item = item_by_slug.get(row["slug"])
        return (
            rank_by_pack.get(item["pack_name_en"], 999) if item else 999,
            tier_order.get(item.get("tier_name_zh"), 9) if item else 9,
            -float(item.get("expected_platinum_contribution") or 0) if item else 0,
            str(row.get("name_zh") or row.get("name_en") or "").casefold(),
        )

    sorted_prices = sorted(prices, key=sort_key)
    headers = [
        "组合包收益排名",
        "洛德组合包",
        "概率池",
        "单次抽中概率(%)",
        "最近日成交数量",
        "计价方式",
        "近48小时加权平均价(满级/白金)",
        "分解荧尘/个",
        "最终计价单价(白金/个)",
        "单包期望贡献(白金)",
        "中文名",
        "期望和",
    ]
    sheet.append(headers)
    _style_header(sheet[1])
    sheet.row_dimensions[1].height = 46

    summary_row_by_pack = {
        row["pack_name_zh"]: 13 + index for index, row in enumerate(summaries)
    }
    groups: list[tuple[str, int, int]] = []
    current_group: str | None = None
    group_start = 2

    for row_number, row in enumerate(sorted_prices, start=2):
        item = item_by_slug.get(row["slug"])
        pack_name = item["pack_name_zh"] if item else "不在洛德组合包"
        if current_group is None:
            current_group = pack_name
            group_start = row_number
        elif pack_name != current_group:
            groups.append((current_group, group_start, row_number - 1))
            current_group = pack_name
            group_start = row_number

        raw_method = item.get("valuation_method") if item else None
        method = str(raw_method or fallback_valuation_method(row, methodology))
        method = "分解再投" if method.startswith("分解再投") and item else method
        if not item and method.startswith("分解再投"):
            method = "无回收价值"

        copies_to_max = (
            int(item["copies_to_max"])
            if item
            else (int(row["max_rank"]) + 1) * (int(row["max_rank"]) + 2) // 2
        )
        sheet.append(
            [
                rank_by_pack.get(item["pack_name_en"], "") if item else "",
                pack_name,
                item.get("tier_name_zh", "") if item else "",
                float(item["probability_per_draw"]) * 100 if item else "",
                int(row.get("latest_daily_volume") or 0),
                method,
                row.get("average_price_48h") if row.get("average_price_48h") is not None else "",
                int(row.get("dissolution_vosfor") or 0),
                None,
                None,
                row.get("name_zh") or row.get("name_en") or row["slug"],
                None,
            ]
        )
        summary_row = summary_row_by_pack.get(pack_name)
        recycle_formula = (
            f"H{row_number}*'{SHEET_PARAMETERS}'!$E${summary_row}/'{SHEET_PARAMETERS}'!$B$5"
            if summary_row
            else "0"
        )
        sheet.cell(row_number, 9).value = (
            f'=IF(F{row_number}="市场出售",G{row_number}/{copies_to_max},'
            f'IF(F{row_number}="分解再投",{recycle_formula},0))'
        )
        if item:
            sheet.cell(row_number, 10).value = f"=3*D{row_number}/100*I{row_number}"
        sheet.cell(row_number, 11).hyperlink = row.get("item_url")
        sheet.cell(row_number, 11).style = "Hyperlink"

    if current_group is not None:
        groups.append((current_group, group_start, len(sorted_prices) + 1))

    for pack_name, start, end in groups:
        sheet.cell(start, 12).value = f"=SUM(J{start}:J{end})"
        fill = GREY_LIGHT if pack_name == "不在洛德组合包" else (
            "F5FAFA" if len([g for g in groups if g[1] < start]) % 2 == 0 else WHITE
        )
        for row in sheet.iter_rows(min_row=start, max_row=end, min_col=1, max_col=12):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=fill)
        sheet.cell(start, 2).fill = PatternFill(
            "solid", fgColor="D9DEE1" if pack_name == "不在洛德组合包" else TEAL_LIGHT
        )
        sheet.cell(start, 2).font = Font(
            name="Microsoft YaHei UI", size=9, bold=True, color="174A50"
        )
        sheet.cell(start, 12).fill = PatternFill(
            "solid", fgColor="E6E9EA" if pack_name == "不在洛德组合包" else "DDF0E6"
        )
        sheet.cell(start, 12).font = Font(
            name="Microsoft YaHei UI", size=10, bold=True, color="145A32"
        )

    min_volume = int(methodology.get("minDailyVolume") or 10)
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=12):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei UI", size=9, color=BODY_TEXT)
            cell.alignment = Alignment(
                horizontal="left" if cell.column == 11 else "center", vertical="center"
            )
            cell.border = _thin_border()
        sheet.row_dimensions[row[0].row].height = 21
        volume_cell = row[4]
        method_cell = row[5]
        if int(volume_cell.value or 0) < min_volume:
            volume_cell.fill = PatternFill("solid", fgColor=RED_LIGHT)
            volume_cell.font = Font(
                name="Microsoft YaHei UI", size=9, bold=True, color=RED_TEXT
            )
        if str(method_cell.value).startswith("分解再投"):
            method_cell.fill = PatternFill("solid", fgColor=YELLOW_LIGHT)
            method_cell.font = Font(
                name="Microsoft YaHei UI", size=9, bold=True, color=YELLOW_TEXT
            )
        elif method_cell.value == "市场出售":
            method_cell.fill = PatternFill("solid", fgColor=GREEN_LIGHT)
            method_cell.font = Font(name="Microsoft YaHei UI", size=9, color=GREEN_TEXT)

    for column in ("D", "G", "I", "J", "L"):
        for cell in sheet[column][1:]:
            cell.number_format = "0.0000"
    for column in ("E", "H"):
        for cell in sheet[column][1:]:
            cell.number_format = "0"
    widths = [11, 24, 12, 15, 14, 12, 20, 13, 21, 18, 24, 13]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    table = Table(displayName="ArcanePriceTable", ref=f"A1:L{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=False, showColumnStripes=False
    )
    sheet.add_table(table)
    sheet["G1"].comment = Comment(
        "价格源：Warframe Market statistics_closed.48hours。以每小时 wa_price 按该小时 volume 再加权。",
        "User",
    )
    return sorted_prices, groups


def _write_parameters_sheet(workbook: Workbook, data: dict[str, Any]) -> None:
    sheet = workbook.create_sheet(SHEET_PARAMETERS)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    methodology = data["methodology"]
    summaries = sorted(data["packSummary"], key=lambda row: int(row["rank"]))
    min_volume = int(methodology.get("minDailyVolume") or 10)
    second_min = int(methodology.get("secondaryFilterMinDailyVolume") or 10)
    second_max = int(methodology.get("secondaryFilterMaxDailyVolume") or 20)
    second_price = float(methodology.get("secondaryFilterMinAveragePrice") or 80)

    sheet.merge_cells("A1:E1")
    sheet["A1"] = "近 48 小时平均价与赋能分解荧尘同包循环再投资模型"
    sheet["A1"].fill = PatternFill("solid", fgColor=TEAL_DARK)
    sheet["A1"].font = Font(name="Microsoft YaHei UI", size=14, bold=True, color=WHITE)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 32

    parameter_rows = [
        ["参数", "值", "说明"],
        ["最近日成交量门槛", min_volume, "低于最近一个已结算自然日门槛时，不按 48 小时平均价出售"],
        ["组合包成本（荧尘）", methodology["packCostVosfor"], "每次购买一个洛德组合包"],
        ["每包抽取数量", methodology["drawsPerPack"], "每个组合包独立抽取 3 个赋能"],
        ["荧尘再投资规则", "购买来源组合包", "每个包分解所得荧尘只继续购买同一个包"],
        ["每 200 荧尘长期价值（白金）", "见下表最终期望/包", "每个组合包分别按无穷等比数列计算"],
        ["每 1 荧尘长期价值（白金）", "各包最终期望 ÷ 200", "明细中的分解再投单价按来源组合包取值"],
        ["追加市场价筛选", f"日量 {second_min}–{second_max} 且满级均价 < {second_price:g}", "满足时剔除市场计价并改按分解荧尘再投资；成交量边界包含，均价等于门槛时保留市场价"],
    ]
    for row_index, values in enumerate(parameter_rows, start=3):
        for column_index, value in enumerate(values, start=1):
            sheet.cell(row_index, column_index).value = value

    summary_header_row = 12
    summary_headers = ["组合包", "直接市场期望", "回收组合包比例", "分解再投贡献", "最终期望/包"]
    for column, value in enumerate(summary_headers, start=1):
        sheet.cell(summary_header_row, column).value = value
    for row_index, summary in enumerate(summaries, start=13):
        sheet.cell(row_index, 1).value = summary["pack_name_zh"]
        sheet.cell(row_index, 2).value = float(summary["direct_market_platinum_per_pack"])
        sheet.cell(row_index, 3).value = float(summary["recycling_pack_fraction"])
        sheet.cell(row_index, 4).value = f"=E{row_index}-B{row_index}"
        sheet.cell(row_index, 5).value = f"=B{row_index}/(1-C{row_index})"

    sheet.merge_cells("A24:C24")
    sheet["A24"] = "计算说明与数据来源"
    notes = [
        ["同包固定点", "V包 = A包 ÷ (1 − B包)", "即 A + BA + B²A + …；后续每轮继续分解并购买同一个包"],
        ["信用点", methodology["packCostCredits"], "未折算成白金；每次用 200 荧尘再买包仍需 50,000 信用点"],
        ["市场数据", "https://api.warframe.market/v1/items/{slug}/statistics", "满级近 48 小时成交量加权平均价；数量门槛用最近已结算自然日"],
        ["荧尘机制", "https://warframe.fandom.com/wiki/Vosfor", "组合包成本与分解机制参考"],
        ["分解数值", methodology["dissolutionDataSource"], "每个未升级赋能的荧尘分解值"],
    ]
    for row_index, values in enumerate(notes, start=25):
        for column_index, value in enumerate(values, start=1):
            sheet.cell(row_index, column_index).value = value

    for header_row, last_column in ((3, 3), (12, 5)):
        for cell in sheet[header_row][:last_column]:
            cell.fill = PatternFill("solid", fgColor=TEAL)
            cell.font = Font(name="Microsoft YaHei UI", size=10, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _thin_border("D6E5E7")
    for row in sheet.iter_rows(min_row=4, max_row=10, min_col=1, max_col=3):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei UI", size=10, color=BODY_TEXT)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = _thin_border()
    for row in sheet.iter_rows(min_row=13, max_row=12 + len(summaries), min_col=1, max_col=5):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei UI", size=10, color=BODY_TEXT)
            cell.alignment = Alignment(vertical="center")
            cell.border = _thin_border()
        for cell in row[1:]:
            cell.number_format = "0.000000" if cell.column == 3 else "0.0000"
    sheet["A24"].fill = PatternFill("solid", fgColor=TEAL_LIGHT)
    sheet["A24"].font = Font(name="Microsoft YaHei UI", size=11, bold=True, color="174A50")
    for row in sheet.iter_rows(min_row=25, max_row=29, min_col=1, max_col=3):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei UI", size=9, color=BODY_TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _thin_border()
    for column, width in {"A": 28, "B": 28, "C": 55, "D": 18, "E": 18}.items():
        sheet.column_dimensions[column].width = width
    for row in range(3, 30):
        sheet.row_dimensions[row].height = 23
    sheet.row_dimensions[10].height = 42
    for row in range(25, 30):
        sheet.row_dimensions[row].height = 38

    sheet["B4"].comment = Comment("低于门槛时改按分解荧尘再投资计价。", "User")
    sheet["B10"].comment = Comment("第二层筛选使用含边界成交量区间；均价等于门槛时保留市场计价。", "User")
    sheet["B7"].comment = Comment("每个组合包独立自循环：V=A+BV=A/(1-B)。", "User")


def validate_workbook(path: Path, data: dict[str, Any]) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    if workbook.sheetnames != [SHEET_DETAIL, SHEET_PARAMETERS]:
        raise RuntimeError(f"工作表不完整：{workbook.sheetnames}")
    formulas = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    if any(token in cell.value for token in ERROR_TOKENS):
                        raise RuntimeError(f"{sheet.title}!{cell.coordinate} 含公式错误标记：{cell.value}")
                    if cell.value.startswith("="):
                        formulas.append(cell.value)
    expected_formula_floor = len(data["prices"]) + len(data["packSummary"]) * 2
    if len(formulas) < expected_formula_floor:
        raise RuntimeError(f"公式数量异常：{len(formulas)} < {expected_formula_floor}")
    workbook.close()


def build_workbook(data: dict[str, Any], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_detail_sheet(workbook, data)
    _write_parameters_sheet(workbook, data)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    validate_workbook(output_path, data)


def main() -> int:
    args = parse_args()
    data = json.loads(args.json_path.read_text(encoding="utf-8"))
    build_workbook(data, args.output_path)
    print(f"Excel：{args.output_path.resolve()}")
    print(f"QA：{len(data['packSummary'])} 个组合包、{len(data['prices'])} 种赋能；工作表与公式检查通过。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

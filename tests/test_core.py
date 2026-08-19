from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import build_arcane_workbook
import warframe_arcane_prices


class CoreMathTests(unittest.TestCase):
    def test_copies_to_max_rank(self) -> None:
        self.assertEqual(warframe_arcane_prices.copies_to_max_rank(5), 21)
        self.assertEqual(warframe_arcane_prices.copies_to_max_rank(3), 10)

    def test_weighted_percentile(self) -> None:
        distribution = [(1.0, 0.25), (2.0, 0.50), (5.0, 0.25)]
        self.assertEqual(warframe_arcane_prices._weighted_percentile(distribution, 0.10), 1.0)
        self.assertEqual(warframe_arcane_prices._weighted_percentile(distribution, 0.50), 2.0)
        self.assertEqual(warframe_arcane_prices._weighted_percentile(distribution, 0.90), 5.0)


class WorkbookTests(unittest.TestCase):
    def test_builds_expected_sheets_and_formulas(self) -> None:
        data = {
            "methodology": {
                "packCostVosfor": 200,
                "packCostCredits": 50000,
                "drawsPerPack": 3,
                "minDailyVolume": 10,
                "secondaryFilterMinDailyVolume": 10,
                "secondaryFilterMaxDailyVolume": 20,
                "secondaryFilterMinAveragePrice": 80,
                "dissolutionDataSource": "https://warframe.fandom.com/wiki/Module:Arcane/data",
            },
            "packSummary": [
                {
                    "rank": 1,
                    "pack_name_zh": "测试赋能组合包",
                    "pack_name_en": "Test Arcane Collection",
                    "direct_market_platinum_per_pack": 4.5,
                    "recycling_pack_fraction": 0.1,
                }
            ],
            "packItems": [
                {
                    "pack_name_zh": "测试赋能组合包",
                    "pack_name_en": "Test Arcane Collection",
                    "tier_name_zh": "稀有",
                    "slug": "arcane_test",
                    "probability_per_draw": 1.0,
                    "copies_to_max": 21,
                    "valuation_method": "市场出售",
                    "expected_platinum_contribution": 4.5,
                }
            ],
            "prices": [
                {
                    "slug": "arcane_test",
                    "name_zh": "测试赋能",
                    "name_en": "Arcane Test",
                    "max_rank": 5,
                    "dissolution_vosfor": 24,
                    "average_price_48h": 31.5,
                    "latest_daily_volume": 100,
                    "item_url": "https://warframe.market/items/arcane_test",
                },
                {
                    "slug": "arcane_outside",
                    "name_zh": "包外赋能",
                    "name_en": "Arcane Outside",
                    "max_rank": 5,
                    "dissolution_vosfor": 0,
                    "average_price_48h": None,
                    "latest_daily_volume": 0,
                    "item_url": "https://warframe.market/items/arcane_outside",
                },
            ],
        }

        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "report.xlsx"
            build_arcane_workbook.build_workbook(data, output)
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                [build_arcane_workbook.SHEET_DETAIL, build_arcane_workbook.SHEET_PARAMETERS],
            )
            detail = workbook[build_arcane_workbook.SHEET_DETAIL]
            parameters = workbook[build_arcane_workbook.SHEET_PARAMETERS]
            self.assertTrue(str(detail["I2"].value).startswith("=IF("))
            self.assertEqual(detail["J2"].value, "=3*D2/100*I2")
            self.assertEqual(parameters["E13"].value, "=B13/(1-C13)")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
